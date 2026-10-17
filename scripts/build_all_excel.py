# -*- coding: utf-8 -*-
"""把 xin 库全部表的数据整理成可读的 Excel（每张表一个工作表 + 总览 + 记账汇总）。"""
import json
import sys
from collections import defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\96082\Desktop\tool\tmp_all.json"
OUT = sys.argv[2] if len(sys.argv) > 2 else r"C:\Users\96082\Desktop\记账数据总表.xlsx"

# ---------------- 样式 ----------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
TOTAL_FONT = Font(bold=True, color="1F4E79")
BAND_FILL = PatternFill("solid", fgColor="F5F8FC")
INCOME_FONT = Font(color="0E7A3D")
EXPENSE_FONT = Font(color="C0392B")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY_FMT = "#,##0.00;[Red]-#,##0.00"
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")

HEADER_MAP = {
    "id": "ID", "user_id": "用户ID", "category_id": "分类ID", "name": "名称",
    "type": "类型", "icon": "图标", "color": "颜色", "sort_order": "排序",
    "is_system": "系统分类", "is_default": "默认分类",
    "created_at": "创建时间", "updated_at": "更新时间", "deleted_at": "删除时间",
    "c_time": "创建时间", "u_time": "更新时间", "d_time": "删除时间", "deleted": "已删除",
    "title": "标题", "remark": "备注", "record_date": "记录日期", "record_time": "记录时间",
    "amount": "金额", "token": "Token", "password": "密码", "salt": "盐值",
    "permission_id": "权限ID", "pid": "父级ID", "creator_id": "创建人ID",
    "description": "描述", "path": "路由", "component": "组件", "hide": "是否隐藏",
    "order_num": "排序", "access": "权限标识", "status": "状态", "role_id": "角色ID",
    "order": "排序", "label": "国际化名称", "department_id": "部门ID",
    "username": "用户名", "email": "邮箱", "phone": "手机号", "is_admin": "账号类型",
    "nick_name": "昵称", "loginip": "登录IP", "login_count": "登录次数",
    "last_login_at": "最后登录时间", "avatar": "头像", "uid": "用户ID",
    "ins": "Instagram", "fb": "Facebook", "weibo": "微博", "zhihu": "知乎",
}

SENSITIVE = {
    ("platform_black_token", "token"),
    ("platform_user", "password"),
    ("platform_user", "salt"),
}

TAB_COLORS = {
    "记账分类": "2E75B6", "记账明细": "1F4E79", "系统分类": "2E75B6", "图标库": "2E75B6",
    "后台用户": "BF8F00", "角色": "BF8F00", "权限": "BF8F00", "角色权限": "BF8F00",
    "用户角色": "BF8F00", "用户联系信息": "BF8F00", "Token黑名单": "BF8F00",
}

data = json.load(open(SRC, encoding="utf-8"))
tables = {t["name"]: t for t in data["tables"]}


def header_for(col):
    name, comment = col["name"], (col.get("comment") or "").strip()
    if name in HEADER_MAP:
        return HEADER_MAP[name]
    return comment or name


def mask(value):
    s = str(value)
    if len(s) <= 8:
        return "***"
    return s[:6] + "…" + s[-4:] + "（已脱敏）"


def decode(table, col, value):
    if value is None:
        return ""
    if col == "type":
        m = {"income": "收入", "expense": "支出", "common": "通用"}
        if value in m:
            return m[value]
    if table == "platform_permission" and col == "type":
        return {0: "目录", 1: "菜单", 2: "按钮"}.get(value, value)
    if col == "is_system" or col == "is_default":
        return "是" if bool(value) else "否"
    if col == "deleted":
        return "是" if bool(value) else "否"
    if col == "hide":
        return "隐藏" if bool(value) else "显示"
    if col == "is_admin":
        return "后台" if bool(value) else "前台"
    if col == "status":
        if table == "platform_user":
            return "启用" if value == 0 else "禁用"
        return "启用" if value == 1 else "禁用"
    return value


def convert(col_type, value):
    if value is None or value == "":
        return ""
    try:
        if col_type == "date":
            return date.fromisoformat(str(value)[:10])
        if col_type in ("datetime", "timestamp"):
            return datetime.fromisoformat(str(value).replace("Z", "")[:19])
        if col_type in ("decimal", "float", "double"):
            return float(value)
        if col_type in ("int", "tinyint", "smallint", "mediumint", "bigint"):
            return int(value)
    except (ValueError, TypeError):
        return value
    return value


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[1].height = 26


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def auto_widths(values_matrix, headers):
    widths = []
    sample = values_matrix[:300]
    for i, h in enumerate(headers):
        best = sum(2 if ord(ch) > 127 else 1 for ch in str(h))
        for row in sample:
            v = row[i]
            if v is None or v == "":
                continue
            s = v.strftime("%Y-%m-%d %H:%M:%S") if isinstance(v, datetime) else (
                v.strftime("%Y-%m-%d") if isinstance(v, date) else str(v))
            best = max(best, min(sum(2 if ord(ch) > 127 else 1 for ch in s), 60))
        widths.append(max(8, min(best + 2, 60)))
    return widths


# ---------------- 总览 ----------------
wb = Workbook()
ws = wb.active
ws.title = "总览"
ws.sheet_properties.tabColor = "1F4E79"
ws.sheet_view.showGridLines = False
set_widths(ws, [7, 36, 24, 15, 10, 10, 10])

ws.merge_cells("A1:G1")
ws["A1"] = "xin 数据库导出总览"
ws["A1"].font = Font(bold=True, size=18, color="1F4E79")
ws["A1"].alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 36

total_rows = sum(len(t["rows"]) for t in data["tables"])
ws["A2"] = f"数据来源：{data['database']} 库 · {len(data['tables'])} 张表 · {total_rows:,} 行数据"
ws["A2"].font = Font(size=11, color="595959")
ws.merge_cells("A2:G2")

headers = ["序号", "表名", "中文说明", "工作表", "行数", "字段数", "打开"]
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = BORDER
ws.row_dimensions[4].height = 24

for i, t in enumerate(data["tables"], start=1):
    r = i + 4
    ws.cell(row=r, column=1, value=i).alignment = CENTER
    ws.cell(row=r, column=2, value=t["name"]).alignment = LEFT
    ws.cell(row=r, column=3, value=t["desc"]).alignment = LEFT
    ws.cell(row=r, column=4, value=t["sheet"]).alignment = CENTER
    ws.cell(row=r, column=5, value=len(t["rows"])).alignment = RIGHT
    ws.cell(row=r, column=6, value=len(t["columns"])).alignment = CENTER
    link = ws.cell(row=r, column=7, value="打开 →")
    link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{t['sheet']}'!A1", display="打开 →")
    link.font = Font(color="0563C1", underline="single")
    link.alignment = CENTER
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = BORDER
    if i % 2 == 0:
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = BAND_FILL

note_row = len(data["tables"]) + 7
ws.cell(row=note_row, column=1, value="说明").font = Font(bold=True, color="404040")
ws.cell(row=note_row, column=2, value=(
    "1. 每张数据库表对应一个工作表，表头取字段注释；日期、金额、枚举值已转为中文可读格式。\n"
    "2. 敏感字段已脱敏：platform_black_token.token、platform_user.password / salt 仅显示首尾少量字符。\n"
    "3. 记账明细共 10,576 条，已包含鲨鱼记账导入数据（2015-02-25 ~ 2026-10-17）。\n"
    "4. 原始数据中有 1,217 条负数支出（退款/冲正），汇总按原值计算，未做调整。"
)).alignment = Alignment(vertical="top")

# ---------------- 各表工作表 ----------------
record_category_names = {}
for tname in ("mini_bookkeeping_category", "mini_bookkeeping_system_category"):
    t = tables.get(tname)
    if not t:
        continue
    names = [c["name"] for c in t["columns"]]
    ni, nn = names.index("id"), names.index("name")
    for row in t["rows"]:
        record_category_names.setdefault(row[ni], row[nn])

for t in data["tables"]:
    name, sheet = t["name"], t["sheet"]
    cols = t["columns"]
    col_names = [c["name"] for c in cols]

    if name == "mini_bookkeeping_record":
        # 记账明细使用可读列布局（保留全部原始字段，另加分类名称/月份）
        idx = {c: i for i, c in enumerate(col_names)}
        headers = ["记录ID", "用户ID", "分类ID", "分类名称", "类型", "金额", "标题", "备注",
                   "记录日期", "月份", "记录时间", "创建时间", "更新时间", "删除时间"]
        matrix = []
        for row in t["rows"]:
            cid = row[idx["category_id"]]
            rdate = str(row[idx["record_date"]] or "")
            rtype = decode(name, "type", row[idx["type"]])
            matrix.append([
                row[idx["id"]], row[idx["user_id"]], cid,
                record_category_names.get(cid, row[idx["title"]] or "未分类"),
                rtype, convert("decimal", row[idx["amount"]]),
                row[idx["title"]] or "", row[idx["remark"]] or "",
                convert("date", rdate), rdate[:7], row[idx["record_time"]] or "",
                convert("timestamp", row[idx["created_at"]]),
                convert("timestamp", row[idx["updated_at"]]),
                convert("datetime", row[idx["deleted_at"]]),
            ])
        aligns = [CENTER, CENTER, CENTER, CENTER, CENTER, RIGHT, LEFT, LEFT,
                  CENTER, CENTER, CENTER, CENTER, CENTER, CENTER]
        money_cols, date_cols = {6}, {9}
        dt_cols = {12, 13, 14}
        widths = [38, 8, 38, 12, 8, 14, 22, 50, 12, 9, 10, 20, 20, 20]
        tab = TAB_COLORS.get(sheet, "808080")
    else:
        headers = [header_for(c) for c in cols]
        matrix = []
        for row in t["rows"]:
            vals = []
            for c, v in zip(cols, row):
                if (name, c["name"]) in SENSITIVE:
                    vals.append(mask(v) if v not in (None, "") else "")
                else:
                    vals.append(decode(name, c["name"], convert(c["type"], v)))
            matrix.append(vals)
        aligns = []
        money_cols, date_cols, dt_cols = set(), set(), set()
        for i, c in enumerate(cols, start=1):
            dt = c["type"]
            if dt in ("int", "tinyint", "smallint", "mediumint", "bigint", "decimal", "float", "double"):
                aligns.append(RIGHT)
            elif dt in ("date", "time", "datetime", "timestamp"):
                aligns.append(CENTER)
            else:
                aligns.append(LEFT)
            if dt == "decimal":
                money_cols.add(i)
            if dt == "date":
                date_cols.add(i)
            if dt in ("datetime", "timestamp"):
                dt_cols.add(i)
        widths = auto_widths(matrix, headers)
        tab = TAB_COLORS.get(sheet, "808080")

    ws = wb.create_sheet(sheet)
    ws.sheet_properties.tabColor = tab
    ws.append(headers)
    style_header(ws, len(headers))

    band = len(matrix) <= 500
    for r_i, values in enumerate(matrix, start=2):
        ws.append(values)
        for c_i in range(1, len(headers) + 1):
            cell = ws.cell(row=r_i, column=c_i)
            cell.alignment = aligns[c_i - 1]
            if band:
                cell.border = BORDER
                if r_i % 2 == 0:
                    cell.fill = BAND_FILL
            if c_i in money_cols:
                cell.number_format = MONEY_FMT
            if c_i in date_cols:
                cell.number_format = "yyyy-mm-dd"
            if c_i in dt_cols:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
        if name == "mini_bookkeeping_record":
            ws.cell(row=r_i, column=5).font = INCOME_FONT if values[4] == "收入" else EXPENSE_FONT
            ws.cell(row=r_i, column=6).font = INCOME_FONT if values[4] == "收入" else EXPENSE_FONT

    set_widths(ws, widths)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

# ---------------- 记账汇总 ----------------
rec = tables["mini_bookkeeping_record"]
rc = {c["name"]: i for i, c in enumerate(rec["columns"])}
records = []
for row in rec["rows"]:
    cid = row[rc["category_id"]]
    records.append({
        "date": str(row[rc["record_date"]]),
        "month": str(row[rc["record_date"]])[:7],
        "type": "收入" if row[rc["type"]] == "income" else "支出",
        "category": record_category_names.get(cid, row[rc["title"]] or "未分类"),
        "amount": float(row[rc["amount"]] or 0),
    })

total_income = round(sum(r["amount"] for r in records if r["type"] == "收入"), 2)
total_expense = round(sum(r["amount"] for r in records if r["type"] == "支出"), 2)
income_cnt = sum(1 for r in records if r["type"] == "收入")
expense_cnt = len(records) - income_cnt
balance = round(total_income - total_expense, 2)

monthly = defaultdict(lambda: {"inc": 0.0, "exp": 0.0, "inc_n": 0, "exp_n": 0})
yearly = defaultdict(lambda: {"inc": 0.0, "exp": 0.0, "inc_n": 0, "exp_n": 0})
categories = defaultdict(lambda: {"inc": 0.0, "exp": 0.0, "inc_n": 0, "exp_n": 0})
for r in records:
    for bucket in (monthly[r["month"]], yearly[r["date"][:4]], categories[r["category"]]):
        if r["type"] == "收入":
            bucket["inc"] += r["amount"]
            bucket["inc_n"] += 1
        else:
            bucket["exp"] += r["amount"]
            bucket["exp_n"] += 1


def add_total_row(ws, values, aligns, money_cols):
    ws.append(values)
    row = ws.max_row
    for c in range(1, len(values) + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.border = BORDER
        cell.alignment = aligns[c - 1]
        if c in money_cols:
            cell.number_format = MONEY_FMT


def build_period_sheet(title, key_label, buckets, tab_color):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    headers = [key_label, "收入（元）", "支出（元）", "结余（元）", "收入笔数", "支出笔数", "合计笔数"]
    ws.append(headers)
    style_header(ws, len(headers))
    aligns = [CENTER, RIGHT, RIGHT, RIGHT, CENTER, CENTER, CENTER]
    for k in sorted(buckets):
        v = buckets[k]
        ws.append([k, round(v["inc"], 2), round(v["exp"], 2), round(v["inc"] - v["exp"], 2),
                   v["inc_n"], v["exp_n"], v["inc_n"] + v["exp_n"]])
        for c in range(1, 8):
            cell = ws.cell(row=ws.max_row, column=c)
            cell.alignment = aligns[c - 1]
            cell.border = BORDER
            if c in (2, 3, 4):
                cell.number_format = MONEY_FMT
    add_total_row(ws, ["合计", total_income, total_expense, balance, income_cnt, expense_cnt, len(records)],
                  aligns, (2, 3, 4))
    set_widths(ws, [11, 15, 15, 15, 10, 10, 10])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row - 1}"
    return ws


build_period_sheet("月度汇总", "月份", monthly, "548235")
year_ws = build_period_sheet("年度汇总", "年份", yearly, "548235")

chart = BarChart()
chart.type = "col"
chart.title = "年度收入 / 支出"
chart.y_axis.title = "金额（元）"
chart.x_axis.title = "年份"
chart.height = 9
chart.width = 20
chart.add_data(Reference(year_ws, min_col=2, max_col=3, min_row=1, max_row=len(yearly) + 1), titles_from_data=True)
chart.set_categories(Reference(year_ws, min_col=1, min_row=2, max_row=len(yearly) + 1))
year_ws.add_chart(chart, "I2")

# 分类汇总
ws = wb.create_sheet("分类汇总")
ws.sheet_properties.tabColor = "548235"
headers = ["分类", "支出金额（元）", "支出笔数", "收入金额（元）", "收入笔数", "净额（元）"]
ws.append(headers)
style_header(ws, len(headers))
aligns = [LEFT, RIGHT, CENTER, RIGHT, CENTER, RIGHT]
items = sorted(categories.items(), key=lambda kv: (kv[1]["exp"], kv[1]["inc"]), reverse=True)
for cat, v in items:
    ws.append([cat, round(v["exp"], 2), v["exp_n"], round(v["inc"], 2), v["inc_n"],
               round(v["inc"] - v["exp"], 2)])
    for c in range(1, 7):
        cell = ws.cell(row=ws.max_row, column=c)
        cell.alignment = aligns[c - 1]
        cell.border = BORDER
        if c in (2, 4, 6):
            cell.number_format = MONEY_FMT
add_total_row(ws, ["合计", total_expense, expense_cnt, total_income, income_cnt, balance],
              aligns, (2, 4, 6))
set_widths(ws, [14, 16, 10, 16, 10, 16])
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:F{ws.max_row - 1}"

chart = BarChart()
chart.type = "bar"
chart.title = "支出分类 TOP 10"
chart.x_axis.title = "金额（元）"
chart.height = 9
chart.width = 18
top_exp = sorted(((k, v["exp"]) for k, v in categories.items() if v["exp_n"]), key=lambda x: x[1], reverse=True)[:10]
chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=len(top_exp) + 1), titles_from_data=False)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(top_exp) + 1))
ws.add_chart(chart, "H2")

wb.properties.title = "xin 数据库导出"
wb.properties.creator = "xin database export"
wb.save(OUT)
print(f"saved {OUT}: {len(data['tables'])} tables, {len(records)} records, {len(wb.sheetnames)} sheets")
